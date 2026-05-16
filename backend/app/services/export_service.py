"""CSV and Excel export helpers."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Iterable

from app.models.cve import CVE
from app.models.profile import OrgProfile


def cves_to_csv(cves: Iterable[CVE]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        [
            "cve_id",
            "severity",
            "cvss_v3_score",
            "cvss_v2_score",
            "is_exploited",
            "kev_vendor",
            "published",
            "last_modified",
            "vendors",
            "cwe",
            "description",
        ]
    )
    for c in cves:
        writer.writerow(
            [
                c.cve_id,
                c.cvss_v3_severity or "",
                c.cvss_v3_score if c.cvss_v3_score is not None else "",
                c.cvss_v2_score if c.cvss_v2_score is not None else "",
                "YES" if c.is_kev else "",
                c.kev_vendor_project or "",
                c.published_date.isoformat() if c.published_date else "",
                c.last_modified_date.isoformat() if c.last_modified_date else "",
                ", ".join(c.vendors or []),
                ", ".join(c.cwe_ids or []),
                (c.description or "").replace("\n", " ").strip(),
            ]
        )
    return buf.getvalue().encode("utf-8")


def cves_to_xlsx(cves: Iterable[CVE]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "CVEs"
    headers = [
        "CVE ID",
        "Severity",
        "CVSS v3",
        "CVSS v2",
        "Exploit Status",
        "Published",
        "Last Modified",
        "Vendors",
        "CWE",
        "Description",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="0F1D3A")
    header_font = Font(bold=True, color="E0A82E")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for c in cves:
        ws.append(
            [
                c.cve_id,
                c.cvss_v3_severity or "",
                c.cvss_v3_score if c.cvss_v3_score is not None else "",
                c.cvss_v2_score if c.cvss_v2_score is not None else "",
                "Exploited" if c.is_kev else "Not exploited",
                c.published_date.isoformat() if c.published_date else "",
                c.last_modified_date.isoformat() if c.last_modified_date else "",
                ", ".join(c.vendors or []),
                ", ".join(c.cwe_ids or []),
                (c.description or "").replace("\n", " ").strip()[:500],
            ]
        )

    widths = [18, 12, 8, 8, 16, 22, 22, 28, 20, 80]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def profile_report_csv(profile: OrgProfile, cves: Iterable[CVE]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["# SentinelX Organisation Vulnerability Report"])
    writer.writerow(["Profile", profile.name])
    writer.writerow(["Asset", profile.asset_name or ""])
    writer.writerow(["Environment", profile.environment.value])
    writer.writerow(["Internet exposed", "YES" if profile.internet_exposed else "NO"])
    writer.writerow(["Business criticality", profile.business_criticality])
    writer.writerow(["Description", profile.description or ""])
    writer.writerow(["Risk Score", profile.risk_score])
    writer.writerow(["Risk Label", profile.risk_label])
    writer.writerow(["Matched CVEs", profile.matched_count])
    writer.writerow(
        [
            "Tech Stack",
            ", ".join(
                f"{e.get('vendor')}:{e.get('product') or '*'}"
                for e in (profile.tech_stack or [])
            ),
        ]
    )
    writer.writerow(["Generated", datetime.now(tz=timezone.utc).isoformat()])
    writer.writerow([])
    writer.writerow(
        [
            "cve_id",
            "cvss_v3_score",
            "cvss_v3_severity",
            "is_kev",
            "vendors",
            "published",
            "description",
        ]
    )
    for c in cves:
        writer.writerow(
            [
                c.cve_id,
                c.cvss_v3_score if c.cvss_v3_score is not None else "",
                c.cvss_v3_severity or "",
                "YES" if c.is_kev else "",
                ", ".join(c.vendors or []),
                c.published_date.isoformat() if c.published_date else "",
                (c.description or "").replace("\n", " ").strip(),
            ]
        )
    return buf.getvalue().encode("utf-8")
